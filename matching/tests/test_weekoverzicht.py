"""Tests for the phase 6 weekoverzicht.

Four things are being protected. First, that a week is assembled from the stored
Tijdblok rows the way the SBTT format expects: one entry per day, totals per
SOORT, a week total. Second, the two instructed differences from the prototype —
no WB column or ⚑ signal anywhere, and a gefactureerd/op-locatie comparison that
really compares Uren.xlsx against the W-blocks. Third, that an incomplete week
says so instead of quietly showing a short week total. Fourth, that the screen
and the Excel export stay two renderings of the same numbers, and that an O
block links to the exact group the uitzonderingenscherm shows.
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import Permission, User
from django.test import TestCase
from django.urls import reverse

from matching import weekoverzicht as week
from matching.models import FaseStatus, LocatieType, Soort, Tijdblok, Uren
from matching.tests import factories
from matching.timeline.runner import run_matching
from matching.weekoverzicht_excel import bestandsnaam, bouw_werkboek

# Week 32 of 2026 runs from Monday 3 to Sunday 9 August — the week the PoC and
# the validation runs both used, so these dates line up with docs/demo.md.
MAANDAG = dt.date(2026, 8, 3)
DINSDAG = dt.date(2026, 8, 4)
WEEK = (2026, 32)

THUIS = ("J. Bosschaartstraat 22", "4112 LN Beusichem")
KLANT = ("Lingedijk 65", "4196 HB Tricht")
VREEMD = ("Industrieweg 2", "4104 AR Culemborg")


def _rit(bestuurder, datum, vertrek, aankomst, van, naar):
    return factories.rit(
        bestuurder,
        datum,
        vertrek,
        aankomst,
        vertrekadres=van[0],
        vertrekplaats=van[1],
        aankomstadres=naar[0],
        aankomstplaats=naar[1],
    )


def _werkdag(bestuurder: str, datum: dt.date) -> None:
    """Home -> customer (2 hours) -> unknown address (1 hour) -> home."""
    _rit(bestuurder, datum, "07:00", "07:30", THUIS, KLANT)
    _rit(bestuurder, datum, "09:30", "09:45", KLANT, VREEMD)
    _rit(bestuurder, datum, "10:45", "11:15", VREEMD, THUIS)


class WeekParsingTests(TestCase):
    """The week value the URL and <input type="week"> both carry."""

    def test_parses_the_iso_week_format(self):
        self.assertEqual(week.parse_week("2026-W32"), (2026, 32))

    def test_rejects_a_week_the_year_does_not_have(self):
        # 2026 has 53 ISO weeks, 2027 does not — the shape alone is not enough.
        self.assertEqual(week.parse_week("2026-W53"), (2026, 53))
        self.assertIsNone(week.parse_week("2027-W53"))

    def test_rejects_nonsense(self):
        for waarde in ("", None, "2026", "week 32", "2026-W", "2026-W99"):
            self.assertIsNone(week.parse_week(waarde), waarde)

    def test_format_round_trips(self):
        self.assertEqual(week.format_week(2026, 32), "2026-W32")
        self.assertEqual(week.parse_week(week.format_week(2026, 7)), (2026, 7))

    def test_week_van_uses_the_iso_year(self):
        # 1 January 2027 is a Friday in week 53 of 2026; the calendar year would
        # point at a week 53 that 2027 does not have.
        self.assertEqual(week.week_van(dt.date(2027, 1, 1)), (2026, 53))


class WeekOpbouwTests(TestCase):
    """§the week itself: days, totals, and the hours comparison."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        _werkdag("M5", MAANDAG)
        _werkdag("M5", DINSDAG)
        factories.urenregel(
            "005",
            MAANDAG,
            "WB260908",
            adres=KLANT[0],
            postcode="4196 HB",
            plaats="TRICHT",
            aantal="2.50",
        )
        # Two lines on one day: the comparison is per day, not per werkbon.
        factories.urenregel(
            "005",
            DINSDAG,
            "WB260908",
            adres=KLANT[0],
            postcode="4196 HB",
            plaats="TRICHT",
            aantal="1.00",
        )
        factories.urenregel(
            "005",
            DINSDAG,
            "WB261000",
            adres=KLANT[0],
            postcode="4196 HB",
            plaats="TRICHT",
            aantal="0.75",
        )
        run_matching()
        self.overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)

    def test_only_days_with_a_timeline_become_a_day(self):
        self.assertEqual(
            [dag.datum for dag in self.overzicht.dagen], [MAANDAG, DINSDAG]
        )

    def test_day_holds_every_block_in_order(self):
        maandag = self.overzicht.dagen[0]
        self.assertEqual(
            [regel.blok.soort for regel in maandag.regels],
            [
                Soort.REISTIJD,
                Soort.WERKBON,
                Soort.REISTIJD,
                Soort.ONVERKLAARD,
                Soort.REISTIJD,
            ],
        )

    def test_soort_totals_cover_every_code_including_the_empty_ones(self):
        totalen = {t.code: t.minuten for t in self.overzicht.dagen[0].totalen}
        self.assertEqual(set(totalen), {s.value for s in Soort})
        self.assertEqual(totalen[Soort.WERKBON], 120)  # 07:30 -> 09:30
        self.assertEqual(totalen[Soort.ONVERKLAARD], 60)  # 09:45 -> 10:45
        self.assertEqual(totalen[Soort.REISTIJD], 30 + 15 + 30)
        self.assertEqual(totalen[Soort.KLANT], 0)

    def test_empty_soort_shows_a_dash_not_a_zero_time(self):
        totalen = {t.code: t.tijd for t in self.overzicht.dagen[0].totalen}
        self.assertEqual(totalen[Soort.WERKBON], "2:00")
        self.assertEqual(totalen[Soort.KLANT], "–")

    def test_day_total_is_every_block_together(self):
        self.assertEqual(self.overzicht.dagen[0].totaal, "4:15")

    def test_week_total_adds_the_days_up(self):
        self.assertEqual(self.overzicht.totaal_minuten, 2 * 255)
        werkbon = next(
            t for t in self.overzicht.totalen if t.code == Soort.WERKBON
        )
        self.assertEqual(werkbon.minuten, 240)

    def test_booked_hours_come_from_uren_per_day(self):
        self.assertEqual(self.overzicht.dagen[0].gefactureerde_uren, Decimal("2.50"))
        # Both lines of Tuesday, added up.
        self.assertEqual(self.overzicht.dagen[1].gefactureerde_uren, Decimal("1.75"))

    def test_hours_on_location_are_the_werkbon_blocks(self):
        # Two hours at the customer; the unexplained hour does not count.
        self.assertEqual(self.overzicht.dagen[0].uren_op_locatie, Decimal("2.00"))
        self.assertEqual(self.overzicht.dagen[0].verschil, Decimal("0.50"))

    def test_week_comparison_totals_both_sides(self):
        self.assertEqual(self.overzicht.gefactureerde_uren, Decimal("4.25"))
        self.assertEqual(self.overzicht.uren_op_locatie, Decimal("4.00"))
        self.assertEqual(self.overzicht.verschil, Decimal("0.25"))

    def test_unexplained_block_carries_the_koppel_key_of_its_group(self):
        onverklaard = next(
            regel
            for regel in self.overzicht.dagen[0].regels
            if regel.blok.soort == Soort.ONVERKLAARD
        )
        self.assertTrue(onverklaard.koppelbaar)
        self.assertEqual(onverklaard.koppel_precisie, LocatieType.POSTCODE)
        self.assertEqual(onverklaard.koppel_sleutel, "4104AR")

    def test_other_blocks_are_not_koppelbaar(self):
        # Only an O is unresolved; offering "koppelen" on a matched werkbon
        # would invite users to link an address that is already explained.
        werkbon = next(
            regel
            for regel in self.overzicht.dagen[0].regels
            if regel.blok.soort == Soort.WERKBON
        )
        self.assertFalse(werkbon.koppelbaar)

    def test_neighbouring_weeks_are_offered(self):
        self.assertEqual(self.overzicht.vorige_week, "2026-W31")
        self.assertEqual(self.overzicht.volgende_week, "2026-W33")

    def test_another_week_is_empty_rather_than_wrong(self):
        leeg = week.bouw_weekoverzicht(self.monteur, 2026, 33)
        self.assertTrue(leeg.leeg)
        self.assertEqual(leeg.totaal_minuten, 0)


class OnvolledigeWeekTests(TestCase):
    """§incomplete week: which days are missing, and never a silent total."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        _werkdag("M5", MAANDAG)
        # Booked hours on a day that has no ride data at all — the case that
        # points at a missing RouteVision export rather than at leave.
        factories.urenregel(
            "005", DINSDAG, "WB260908", adres=KLANT[0], postcode="4196 HB", aantal="8.00"
        )
        run_matching()

    def test_missing_working_days_are_reported(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        self.assertEqual(
            [dag.datum for dag in overzicht.ontbrekende_dagen],
            [
                dt.date(2026, 8, 4),
                dt.date(2026, 8, 5),
                dt.date(2026, 8, 6),
                dt.date(2026, 8, 7),
            ],
        )

    def test_weekend_is_not_missing(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        datums = [dag.datum for dag in overzicht.ontbrekende_dagen]
        self.assertNotIn(dt.date(2026, 8, 8), datums)
        self.assertNotIn(dt.date(2026, 8, 9), datums)

    def test_missing_day_carries_the_hours_booked_on_it(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        dinsdag = overzicht.ontbrekende_dagen[0]
        self.assertEqual(dinsdag.gefactureerde_uren, Decimal("8.00"))

    def test_hours_of_a_missing_day_stay_out_of_the_week_total(self):
        # Counting Tuesday's 8 hours against a Tuesday that has no W-blocks
        # would read as a discrepancy in the work, while it is a gap in the data.
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        self.assertEqual(overzicht.gefactureerde_uren, Decimal("0.00"))

    def test_future_days_are_not_missing(self):
        # The current week, halfway through: the days still to come are not a
        # gap. Nothing is processed for this monteur in it, so every day that
        # has already passed is missing and none of the later ones are.
        vandaag = dt.date.today()
        overzicht = week.bouw_weekoverzicht(self.monteur, *week.week_van(vandaag))
        for dag in overzicht.ontbrekende_dagen:
            self.assertLessEqual(dag.datum, vandaag)


class WeekoverzichtSchermTests(TestCase):
    """§the page: navigation, the two dropped columns, and the O link."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        self.collega = factories.monteur("Dennis", "001", "M1")
        _werkdag("M5", MAANDAG)
        _werkdag("M1", MAANDAG)
        factories.urenregel(
            "005", MAANDAG, "WB260908", adres=KLANT[0], postcode="4196 HB", aantal="2.50"
        )
        run_matching()

        self.gebruiker = User.objects.create_user("wim", password="geheim")
        self.gebruiker.user_permissions.add(
            Permission.objects.get(codename="add_bekendelocatie")
        )
        self.client.force_login(self.gebruiker)
        self.url = reverse("weekoverzicht")

    def _pagina(self, **params) -> str:
        antwoord = self.client.get(self.url, params)
        self.assertEqual(antwoord.status_code, 200)
        return antwoord.content.decode()

    def test_login_is_required(self):
        self.client.logout()
        antwoord = self.client.get(self.url)
        self.assertEqual(antwoord.status_code, 302)
        self.assertIn("/admin/login/", antwoord["Location"])

    def test_defaults_to_the_last_processed_week(self):
        # No parameters at all: the screen opens on a week that has content
        # rather than on an empty current week. This is the default the root
        # URL and the navigation link both land on, so it must never be blank
        # (see week.laatste_week_met_data for why it is not "this week").
        inhoud = self._pagina()
        self.assertIn("2026-W32", inhoud)
        self.assertIn("Jesse", inhoud)

    def test_defaults_to_the_first_monteur_alphabetically(self):
        # Dennis before Jesse, whatever order they were created in — "the first
        # one alphabetically" is the defined default of a bare /weekoverzicht/.
        self.assertIn("Weekoverzicht — Dennis", self._pagina())

    def test_the_default_page_is_not_empty(self):
        # The whole point of the two defaults above: someone arriving through /
        # or through the navigation bar sees a filled week, not a blank shell
        # with two dropdowns.
        inhoud = self._pagina()
        self.assertNotIn("is nog geen dag berekend", inhoud)
        self.assertIn("Maandag 3 augustus", inhoud)

    def test_selects_the_monteur_from_the_url(self):
        inhoud = self._pagina(monteur=self.collega.pk, week="2026-W32")
        self.assertIn("Weekoverzicht — Dennis", inhoud)

    def test_unknown_monteur_is_a_404_not_a_silent_fallback(self):
        antwoord = self.client.get(self.url, {"monteur": 9999})
        self.assertEqual(antwoord.status_code, 404)

    def test_impossible_week_is_a_404(self):
        antwoord = self.client.get(self.url, {"week": "2027-W53"})
        self.assertEqual(antwoord.status_code, 404)

    def test_no_wb_column_and_no_flag_signal(self):
        # The instructed omission: the WB-vs-SYS comparison is not built, so it
        # must not reappear here through the prototype's layout.
        inhoud = self._pagina(week="2026-W32")
        self.assertNotIn("⚑", inhoud)
        self.assertNotIn("<th>WB</th>", inhoud)
        self.assertNotIn("SYS aank", inhoud)

    def test_shows_the_hours_comparison(self):
        # "Totaal (excl. reistijd)" rather than "Gefactureerd": depot and
        # magazijn time is booked but never invoiced, so the old label promised
        # more than the number meant (docs/decisions.md, 07-09-2026 avond).
        inhoud = self._pagina(week="2026-W32")
        self.assertIn("Totaal (excl. reistijd)", inhoud)
        self.assertIn("op locatie", inhoud)
        self.assertNotIn("Gefactureerd", inhoud)

    def test_unexplained_block_links_to_the_koppel_form(self):
        inhoud = self._pagina(week="2026-W32")
        verwacht = reverse(
            "uitzonderingen_koppelen", args=[LocatieType.POSTCODE, "4104AR"]
        )
        self.assertIn(verwacht, inhoud)

    def test_koppel_link_hidden_without_the_permission(self):
        kijker = User.objects.create_user("kijker", password="geheim")
        self.client.force_login(kijker)
        inhoud = self._pagina(week="2026-W32")
        self.assertNotIn("koppelen/", inhoud)

    def test_koppel_link_opens_the_group_the_list_shows(self):
        # The link is only useful if it lands on a real group; phase 5 answers
        # a 404 for a group that is not in the list.
        doel = reverse("uitzonderingen_koppelen", args=[LocatieType.POSTCODE, "4104AR"])
        self.assertEqual(self.client.get(doel).status_code, 200)

    def test_incomplete_week_warns_on_screen(self):
        inhoud = self._pagina(week="2026-W32")
        self.assertIn("Onvolledige week", inhoud)

    def test_empty_week_says_so_instead_of_showing_zero_totals(self):
        inhoud = self._pagina(week="2026-W20")
        self.assertIn("nog geen dag berekend", inhoud)


class WeekoverzichtExcelTests(TestCase):
    """§the export: same numbers, same colours, as a real workbook."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse Verkerk", "005", "M5")
        _werkdag("M5", MAANDAG)
        factories.urenregel(
            "005", MAANDAG, "WB260908", adres=KLANT[0], postcode="4196 HB", aantal="2.50"
        )
        run_matching()
        self.overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)

    def _blad(self):
        werkboek = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(self.overzicht)))
        return werkboek[werkboek.sheetnames[0]]

    def _cellen(self, blad) -> list[str]:
        return [
            str(cel.value)
            for rij in blad.iter_rows()
            for cel in rij
            if cel.value is not None
        ]

    def test_filename_names_monteur_and_week(self):
        self.assertEqual(
            bestandsnaam(self.overzicht),
            "RMW-weekoverzicht-2026-W32-Jesse-Verkerk.xlsx",
        )

    def test_sheet_holds_the_days_and_their_blocks(self):
        cellen = self._cellen(self._blad())
        self.assertIn("Maandag 03-08-2026 · totaal 4:15", cellen)
        self.assertIn("WB260908", cellen)
        self.assertIn("Onverklaarde stop", cellen)

    def test_soort_cells_carry_the_screen_colour(self):
        blad = self._blad()
        gekleurd = {
            cel.value: cel.fill.fgColor.rgb
            for rij in blad.iter_rows()
            for cel in rij
            if cel.value in {Soort.WERKBON.value, Soort.ONVERKLAARD.value}
            and cel.fill.fgColor.rgb not in (None, "00000000")
        }
        # openpyxl prefixes the alpha channel; the colour itself is the one the
        # page paints the pill with.
        self.assertTrue(gekleurd[Soort.WERKBON].endswith("1F8A4C"))
        self.assertTrue(gekleurd[Soort.ONVERKLAARD].endswith("D97706"))

    def test_hours_are_written_as_numbers_not_as_text(self):
        blad = self._blad()
        # A whole number of hours comes back as an int, so both types count:
        # what matters is that the cell holds a number a user can total further,
        # not the text "2,00 u".
        getallen = [
            cel.value
            for rij in blad.iter_rows()
            for cel in rij
            if isinstance(cel.value, (int, float)) and not isinstance(cel.value, bool)
        ]
        self.assertIn(2.5, getallen)  # gefactureerd
        self.assertIn(2, getallen)  # op locatie
        self.assertIn(0.5, getallen)  # verschil

    def test_the_hours_comparison_carries_the_same_label_as_the_page(self):
        # The export and the page must not disagree about what the number is
        # called (docs/decisions.md, 07-09-2026 avond).
        cellen = self._cellen(self._blad())
        self.assertIn("Totaal (excl. reistijd) (Uren.xlsx)", cellen)
        self.assertNotIn("Gefactureerde uren (Uren.xlsx)", cellen)

    def test_no_wb_column_and_no_flag_signal(self):
        cellen = self._cellen(self._blad())
        self.assertNotIn("WB", cellen)
        self.assertNotIn("⚑", cellen)

    def test_incomplete_week_is_stated_in_the_sheet(self):
        cellen = " ".join(self._cellen(self._blad()))
        self.assertIn("onvolledige week", cellen.lower())

    def test_download_serves_a_workbook(self):
        gebruiker = User.objects.create_user("wim", password="geheim")
        self.client.force_login(gebruiker)
        antwoord = self.client.get(
            reverse("weekoverzicht_excel"),
            {"monteur": self.monteur.pk, "week": "2026-W32"},
        )
        self.assertEqual(antwoord.status_code, 200)
        self.assertIn("spreadsheetml", antwoord["Content-Type"])
        self.assertIn(
            "RMW-weekoverzicht-2026-W32-Jesse-Verkerk.xlsx",
            antwoord["Content-Disposition"],
        )
        # A real workbook, not an error page with the right header on it.
        openpyxl.load_workbook(io.BytesIO(antwoord.content))

    def test_download_requires_login(self):
        antwoord = self.client.get(reverse("weekoverzicht_excel"))
        self.assertEqual(antwoord.status_code, 302)


class LegeInrichtingTests(TestCase):
    """The state before anything is set up: no monteurs at all."""

    def setUp(self):
        gebruiker = User.objects.create_user("wim", password="geheim")
        self.client.force_login(gebruiker)

    def test_page_explains_what_is_missing(self):
        antwoord = self.client.get(reverse("weekoverzicht"))
        self.assertEqual(antwoord.status_code, 200)
        self.assertIn("nog geen monteurs", antwoord.content.decode())

    def test_export_has_nothing_to_export(self):
        antwoord = self.client.get(reverse("weekoverzicht_excel"))
        self.assertEqual(antwoord.status_code, 404)


class PriveTests(TestCase):
    """SOORT P: an address someone marked private (07-09-2026).

    P is a hand-assigned classification like K/L/C, so it needs no new rule in
    the engine — it rides along in the "rest of the koppeltabel" step. What is
    new is that its time is reported on its own: not as work, and not as a
    deduction from anything either.
    """

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        _werkdag("M5", MAANDAG)
        factories.urenregel(
            "005", MAANDAG, "WB260908",
            adres=KLANT[0], postcode="4196 HB", plaats="TRICHT", aantal="2.50",
        )

    def _overzicht(self):
        run_matching(force=True)
        return week.bouw_weekoverzicht(self.monteur, *WEEK)

    def _markeer_vreemd_adres_als_prive(self):
        # VREEMD is the hour-long stop of _werkdag that would otherwise be O.
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AR", Soort.PRIVE, "Sportschool"
        )

    def test_an_address_linked_as_prive_classifies_its_stop_as_p(self):
        self._markeer_vreemd_adres_als_prive()
        dag = self._overzicht().dagen[0]

        soorten = [regel.blok.soort for regel in dag.regels]
        self.assertIn(Soort.PRIVE, soorten)
        self.assertNotIn(Soort.ONVERKLAARD, soorten)

    def test_prive_time_is_reported_as_its_own_figure(self):
        self._markeer_vreemd_adres_als_prive()
        overzicht = self._overzicht()

        # The VREEMD stop lasts an hour in _werkdag.
        self.assertEqual(overzicht.dagen[0].prive_uren, Decimal("1.00"))
        self.assertEqual(overzicht.prive_uren, Decimal("1.00"))

    def test_prive_time_does_not_touch_the_hours_comparison(self):
        zonder = self._overzicht()
        geboekt, op_locatie = zonder.gefactureerde_uren, zonder.uren_op_locatie

        self._markeer_vreemd_adres_als_prive()
        met = self._overzicht()

        # Marking an address private must not move the work figures a millimetre
        # in either direction — it is a number beside them, not a correction.
        self.assertEqual(met.gefactureerde_uren, geboekt)
        self.assertEqual(met.uren_op_locatie, op_locatie)
        self.assertEqual(met.verschil, zonder.verschil)

    def test_a_week_without_prive_stops_reports_zero(self):
        overzicht = self._overzicht()
        self.assertEqual(overzicht.prive_uren, Decimal("0.00"))

    def test_p_has_its_own_colour_distinct_from_the_other_seven(self):
        kleuren = week.SOORT_KLEUREN
        self.assertIn(Soort.PRIVE, kleuren)
        self.assertEqual(len(set(kleuren.values())), len(Soort))

    def test_p_is_totalled_with_the_other_soorten(self):
        self._markeer_vreemd_adres_als_prive()
        overzicht = self._overzicht()

        codes = [totaal.code for totaal in overzicht.totalen]
        self.assertIn(Soort.PRIVE.value, codes)
        prive = next(t for t in overzicht.totalen if t.code == Soort.PRIVE.value)
        self.assertEqual(prive.minuten, 60)


class PriveWeergaveTests(TestCase):
    """The privé figure has to reach both renderings, page and Excel."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AR", Soort.PRIVE, "Sportschool"
        )
        _werkdag("M5", MAANDAG)
        run_matching(force=True)

        self.gebruiker = User.objects.create_user("kijker", "k@sbtt.nl", "geheim")
        self.client.force_login(self.gebruiker)

    def _pagina(self):
        antwoord = self.client.get(
            reverse("weekoverzicht"),
            {"monteur": self.monteur.pk, "week": "2026-W32"},
        )
        return antwoord.content.decode()

    def test_the_page_shows_the_prive_figure(self):
        inhoud = self._pagina()
        self.assertIn("Privé", inhoud)
        self.assertIn("#C2185B", inhoud)

    def test_the_page_says_it_is_not_deducted(self):
        # The wording matters: without it the number reads as a correction on
        # the totals above it.
        self.assertIn("niet van de", self._pagina())

    def test_the_excel_export_carries_the_same_figure(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        boek = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(overzicht)))
        blad = boek.active
        cellen = [
            cel.value for rij in blad.iter_rows() for cel in rij if cel.value is not None
        ]

        self.assertIn("Privé (SOORT P)", cellen)
        self.assertIn(1, cellen)  # one hour privé, written as a number


class ThuisWeergaveTests(TestCase):
    """SOORT T: a code like any other in the table, with no summary of its own.

    Deliberately unlike P (07-09-2026, docs/decisions.md): time at home is less
    of a number you want totalled than privé time is, so T gets a colour and a
    row in the day table and nothing beside the week figures.
    """

    def setUp(self):
        self.monteur = factories.monteur(
            "Jesse", "005", "M5", thuisadres=THUIS[0]
        )
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        # Home -> customer -> home, so the day has a stop at home in the middle.
        _rit("M5", MAANDAG, "07:00", "07:30", THUIS, KLANT)
        _rit("M5", MAANDAG, "09:30", "10:00", KLANT, THUIS)
        _rit("M5", MAANDAG, "12:00", "12:30", THUIS, KLANT)
        _rit("M5", MAANDAG, "16:00", "16:30", KLANT, THUIS)
        run_matching(force=True)

        self.gebruiker = User.objects.create_user("kijker", "k@sbtt.nl", "geheim")
        self.client.force_login(self.gebruiker)

    def _pagina(self):
        return self.client.get(
            reverse("weekoverzicht"),
            {"monteur": self.monteur.pk, "week": "2026-W32"},
        ).content.decode()

    def test_the_stop_at_home_is_classified_as_thuis(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        soorten = [regel.blok.soort for regel in overzicht.dagen[0].regels]

        self.assertIn(Soort.THUIS, soorten)

    def test_t_has_a_colour_of_its_own(self):
        self.assertIn(Soort.THUIS, week.SOORT_KLEUREN)
        self.assertEqual(len(set(week.SOORT_KLEUREN.values())), len(Soort))
        self.assertIn(week.SOORT_KLEUREN[Soort.THUIS], self._pagina())

    def test_t_is_totalled_per_soort_like_every_other_code(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        codes = [totaal.code for totaal in overzicht.totalen]

        self.assertIn(Soort.THUIS.value, codes)

    def test_t_gets_no_summary_line_of_its_own(self):
        # P has one ("Privé ... uur"); T deliberately does not.
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)

        self.assertFalse(hasattr(overzicht, "thuis_uren"))
        self.assertNotIn("Thuis 2", self._pagina())

    def test_t_does_not_touch_the_hours_comparison(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)

        # Only the customer stops count as time on location; home never does.
        self.assertEqual(overzicht.uren_op_locatie, Decimal("0.00"))
        self.assertEqual(overzicht.prive_uren, Decimal("0.00"))

    def test_the_excel_export_has_no_thuis_summary_row_either(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        boek = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(overzicht)))
        cellen = [
            cel.value
            for rij in boek.active.iter_rows()
            for cel in rij
            if cel.value is not None
        ]

        self.assertIn("Thuis", cellen)  # the per-SOORT total row
        self.assertNotIn("Thuis (SOORT T)", cellen)  # but no comparison row


class AansluitingPerWerkbonTests(TestCase):
    """The per-werkbon reconciliation under each day (07-09-2026).

    The day comparison above it blends every werkbon into one pair of numbers,
    so it cannot show time spent on werkbon A while the hours went onto werkbon
    B. These tests pin down that every werkbon of the day gets a row from either
    source, that a side which cannot exist reads as absent rather than as zero,
    and that the total is the same number the day line already shows.
    """

    TWEEDE_KLANT = ("Molenstraat 5", "4191 AA Geldermalsen")

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        # Two customer stops of an hour each, on two different werkbonnen.
        _rit("M5", MAANDAG, "07:00", "07:30", THUIS, KLANT)
        _rit("M5", MAANDAG, "08:30", "09:00", KLANT, self.TWEEDE_KLANT)
        _rit("M5", MAANDAG, "10:00", "10:30", self.TWEEDE_KLANT, THUIS)
        factories.urenregel(
            "005", MAANDAG, "WB1000",
            adres=KLANT[0], postcode="4196 HB", plaats="TRICHT", aantal="1.00",
            opdrachtgever="Bakkerij Van Dijk",
        )
        factories.urenregel(
            "005", MAANDAG, "WB2000",
            adres=self.TWEEDE_KLANT[0], postcode="4191 AA", plaats="GELDERMALSEN",
            aantal="1.00", opdrachtgever="Garage Molenstraat",
        )

    def _dag(self):
        run_matching(force=True)
        return week.bouw_weekoverzicht(self.monteur, *WEEK).dagen[0]

    def _regels(self, dag):
        return {regel.label: regel for regel in dag.aansluiting}

    def _tweede_stop_vrijgeven(self):
        """Drop the booked hours of the second stop, so it is not a W any more.

        Booked hours outrank the koppeltabel in the priority order, so a stop
        with an Uren row on it can never come out as K, C or a Werkbonnen.xlsx
        fallback — the classification has to be freed up first.
        """
        Uren.objects.filter(werkbon="WB2000").delete()

    def test_each_werkbon_of_the_day_gets_its_own_row(self):
        regels = self._regels(self._dag())

        self.assertIn("WB1000", regels)
        self.assertIn("WB2000", regels)
        self.assertEqual(regels["WB1000"].op_locatie, Decimal("1.00"))
        self.assertEqual(regels["WB1000"].gedeclareerd, Decimal("1.00"))
        self.assertEqual(regels["WB1000"].verschil, Decimal("0.00"))
        self.assertTrue(regels["WB1000"].sluit_aan)

    def test_hours_booked_on_a_werkbon_with_no_time_on_site_still_get_a_row(self):
        # The loudest signal the table can give, so it must never be the row
        # that happens to be missing.
        factories.urenregel("005", MAANDAG, "WB9999", aantal="4.00")

        regel = self._regels(self._dag())["WB9999"]

        self.assertEqual(regel.op_locatie, Decimal("0.00"))
        self.assertEqual(regel.gedeclareerd, Decimal("4.00"))
        self.assertEqual(regel.verschil, Decimal("4.00"))
        self.assertFalse(regel.sluit_aan)

    def test_time_on_site_without_any_booked_hours_also_gets_a_row(self):
        # The reverse case: an hour spent on a werkbon nobody declared. It
        # reaches the timeline through the Werkbonnen.xlsx postcode fallback,
        # which is the one path that produces a W block without an Uren row.
        self._tweede_stop_vrijgeven()
        factories.werkbon_controle(
            "WB2000", "005", MAANDAG, FaseStatus.AFGEROND, postcode="4191 AA"
        )

        regel = self._regels(self._dag())["WB2000"]

        self.assertEqual(regel.op_locatie, Decimal("1.00"))
        self.assertEqual(regel.gedeclareerd, Decimal("0.00"))
        self.assertEqual(regel.verschil, Decimal("-1.00"))

    def test_hours_without_a_werkbon_number_get_their_own_row(self):
        # Kantoor, verlof, reisuren: real booked hours that can never produce a
        # W block. Leaving them out would make the table total disagree with the
        # day figure right above it.
        factories.urenregel("005", MAANDAG, "", aantal="2.00")

        regel = self._regels(self._dag())[week.INDIRECT_LABEL]

        self.assertIsNone(regel.op_locatie)
        self.assertEqual(regel.gedeclareerd, Decimal("2.00"))
        self.assertIsNone(regel.verschil)

    def test_there_is_no_indirect_row_when_every_line_has_a_werkbon(self):
        self.assertNotIn(week.INDIRECT_LABEL, self._regels(self._dag()))

    def test_client_time_is_its_own_row_without_a_declared_figure(self):
        # A BekendeLocatie(K) carries no werkbon number anywhere in the data, so
        # there is nothing to compare it against — absent, not zero.
        self._tweede_stop_vrijgeven()
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4191 AA", Soort.KLANT, "Vaste klant"
        )

        regel = self._regels(self._dag())[week.KLANT_LABEL]

        self.assertEqual(regel.op_locatie, Decimal("1.00"))
        self.assertIsNone(regel.gedeclareerd)
        self.assertIsNone(regel.verschil)

    def test_the_client_row_is_shown_even_when_there_is_no_client_time(self):
        # An absent row would read as "no client time unaccounted for", which is
        # exactly the confusion this table exists to remove.
        regel = self._regels(self._dag())[week.KLANT_LABEL]

        self.assertEqual(regel.op_locatie, Decimal("0.00"))

    def test_client_time_counts_towards_the_day_total(self):
        self._tweede_stop_vrijgeven()
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4191 AA", Soort.KLANT, "Vaste klant"
        )
        dag = self._dag()

        # One hour of W (WB1000) plus one hour of K.
        self.assertEqual(dag.aansluiting_totaal.op_locatie, Decimal("2.00"))

    def test_locatie_and_crediteur_stay_out_of_the_total(self):
        # Deliberately client-linked hours only (docs/decisions.md, 07-09-2026).
        self._tweede_stop_vrijgeven()
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4191 AA", Soort.CREDITEUR, "Groothandel"
        )
        dag = self._dag()

        self.assertEqual(dag.aansluiting_totaal.op_locatie, Decimal("1.00"))

    def test_the_total_matches_the_day_comparison_above_it(self):
        # Two totals a few cents apart would cost more trust than the cents are
        # worth, so the declared side is the day figure itself.
        dag = self._dag()

        self.assertEqual(dag.aansluiting_totaal.gedeclareerd, dag.gefactureerde_uren)
        self.assertEqual(dag.aansluiting_totaal.op_locatie, dag.uren_op_locatie)

    def test_a_werkbon_row_shows_the_client_it_was_booked_for(self):
        # The reason for the column (docs/decisions.md, 08-09-2026): judging a
        # difference should not need a lookup elsewhere to see whose job it was.
        regels = self._regels(self._dag())

        self.assertEqual(regels["WB1000"].klantnaam, "Bakkerij Van Dijk")
        self.assertEqual(regels["WB2000"].klantnaam, "Garage Molenstraat")

    def test_a_werkbon_without_an_uren_row_that_day_shows_no_name(self):
        # Reached through the Werkbonnen.xlsx postcode fallback, so there is no
        # Uren row on this date to read a name from. Absent rather than borrowed
        # from another date of the same werkbon.
        self._tweede_stop_vrijgeven()
        factories.werkbon_controle(
            "WB2000", "005", MAANDAG, FaseStatus.AFGEROND, postcode="4191 AA"
        )

        self.assertIsNone(self._regels(self._dag())["WB2000"].klantnaam)

    def test_the_name_is_not_looked_up_on_another_date_of_the_same_werkbon(self):
        # Same werkbon, named on Tuesday only: Monday still shows a dash. The
        # column says what stands on this exact (date, werkbon), nothing wider.
        self._tweede_stop_vrijgeven()
        factories.werkbon_controle(
            "WB2000", "005", MAANDAG, FaseStatus.AFGEROND, postcode="4191 AA"
        )
        factories.urenregel(
            "005", MAANDAG + dt.timedelta(days=1), "WB2000",
            aantal="1.00", opdrachtgever="Garage Molenstraat",
        )

        self.assertIsNone(self._regels(self._dag())["WB2000"].klantnaam)

    def test_the_indirect_and_client_rows_carry_no_name(self):
        # Neither is tied to one werkbon, so neither can name a client.
        factories.urenregel(
            "005", MAANDAG, "", aantal="2.00", opdrachtgever="Kantoor"
        )
        dag = self._dag()
        regels = self._regels(dag)

        self.assertIsNone(regels[week.INDIRECT_LABEL].klantnaam)
        self.assertIsNone(regels[week.KLANT_LABEL].klantnaam)
        self.assertIsNone(dag.aansluiting_totaal.klantnaam)

    def test_a_day_with_nothing_to_reconcile_shows_no_table(self):
        self._dag()
        Uren.objects.all().delete()
        Tijdblok.objects.filter(soort=Soort.WERKBON).delete()

        dag = week.bouw_weekoverzicht(self.monteur, *WEEK).dagen[0]

        self.assertFalse(dag.heeft_aansluiting)


class AansluitingWeergaveTests(TestCase):
    """The table has to reach both renderings, page and Excel."""

    def setUp(self):
        self.monteur = factories.monteur("Jesse", "005", "M5")
        factories.bekende_locatie(
            LocatieType.POSTCODE, "4104 AC", Soort.LOCATIE, "Magazijn", is_depot=True
        )
        _werkdag("M5", MAANDAG)
        factories.urenregel(
            "005", MAANDAG, "WB260908",
            adres=KLANT[0], postcode="4196 HB", plaats="TRICHT", aantal="2.50",
            opdrachtgever="Bakkerij Van Dijk",
        )
        # A werkbon with hours but no time on site: the mismatch to look for.
        # Booked without a client name, so it renders as a dash.
        factories.urenregel("005", MAANDAG, "WB260999", aantal="4.00")
        run_matching(force=True)

        self.gebruiker = User.objects.create_user("kijker", "k@sbtt.nl", "geheim")
        self.client.force_login(self.gebruiker)

    def _pagina(self):
        return self.client.get(
            reverse("weekoverzicht"),
            {"monteur": self.monteur.pk, "week": "2026-W32"},
        ).content.decode()

    def test_the_page_shows_the_table_with_both_werkbonnen(self):
        inhoud = self._pagina()

        self.assertIn("Aansluiting per werkbon", inhoud)
        self.assertIn("Op locatie", inhoud)
        self.assertIn("Gedeclareerd", inhoud)
        self.assertIn("WB260908", inhoud)
        self.assertIn("WB260999", inhoud)

    def test_the_page_keeps_the_existing_day_comparison_as_well(self):
        # The new table sits alongside the old line, not instead of it.
        inhoud = self._pagina()

        self.assertIn("Totaal (excl. reistijd)", inhoud)
        self.assertIn("Aansluiting per werkbon", inhoud)

    def test_the_page_shows_the_client_row(self):
        self.assertIn(week.KLANT_LABEL, self._pagina())

    def test_the_page_shows_the_client_column(self):
        inhoud = self._pagina()

        self.assertIn("<th>Klant</th>", inhoud)
        self.assertIn("Bakkerij Van Dijk", inhoud)

    def test_the_excel_export_carries_the_same_table(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        boek = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(overzicht)))
        cellen = [
            cel.value
            for rij in boek.active.iter_rows()
            for cel in rij
            if cel.value is not None
        ]

        self.assertIn("Aansluiting per werkbon", cellen)
        self.assertIn("WB260908", cellen)
        self.assertIn("WB260999", cellen)
        self.assertIn(week.KLANT_LABEL, cellen)
        self.assertIn(4.0, cellen)  # the undeclared werkbon gap, as a number

    def test_the_excel_export_carries_the_client_column(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        blad = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(overzicht))).active
        rijen = {r[0].value: r for r in blad.iter_rows()}

        self.assertEqual(rijen["Werkbon"][1].value, "Klant")
        self.assertEqual(rijen["WB260908"][1].value, "Bakkerij Van Dijk")
        # No name on this date, and the rows that are not tied to one werkbon:
        # a dash, the same "not applicable" the page shows.
        self.assertEqual(rijen["WB260999"][1].value, "–")
        self.assertEqual(rijen[week.KLANT_LABEL][1].value, "–")
        self.assertEqual(rijen["Totaal"][1].value, "–")

    def test_the_excel_export_leaves_an_absent_side_empty(self):
        overzicht = week.bouw_weekoverzicht(self.monteur, *WEEK)
        boek = openpyxl.load_workbook(io.BytesIO(bouw_werkboek(overzicht)))

        rij = next(
            r for r in boek.active.iter_rows() if r[0].value == week.KLANT_LABEL
        )
        # Op locatie is written; gedeclareerd and verschil stay empty, because
        # the question does not apply rather than having come back zero.
        self.assertEqual(rij[2].value, 0.0)
        self.assertIsNone(rij[3].value)
        self.assertIsNone(rij[4].value)
